from flask import Flask, request
import openpyxl
from datetime import datetime
import os

app = Flask(__name__)

# Detect the exact folder directory where app.py is located
BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# Define absolute file paths so Python never loses track of the Excel files
ATTENDANCE_FILE = os.path.join(BASE_DIR, "Attendance.xlsx")
STUDENTS_FILE = os.path.join(BASE_DIR, "Students.xlsx")

last_scanned_times = {}

def initialize_files():
    """Automatically creates Attendance.xlsx and Students.xlsx if they don't exist."""
    if not os.path.exists(ATTENDANCE_FILE):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Attendance Log"
        ws.append(["Timestamp", "UID", "Name", "Status"])
        wb.save(ATTENDANCE_FILE)
        print(f"[INIT] Created {ATTENDANCE_FILE}")

    if not os.path.exists(STUDENTS_FILE):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Students List"
        ws.append(["UID", "Name", "Registered Date"])
        wb.save(STUDENTS_FILE)
        print(f"[INIT] Created {STUDENTS_FILE}")

@app.route('/log', methods=['GET'])
def log_attendance():
    uid = request.args.get('uid')
    name = request.args.get('name')

    if not uid or not name:
        return "MISSING_PARAMS", 400

    now = datetime.now()
    timestamp_str = now.strftime("%Y-%m-%d %H:%M:%S")

    # Prevent duplicate scans within 5 minutes (300 seconds)
    if uid in last_scanned_times:
        time_diff = (now - last_scanned_times[uid]).total_seconds()
        if time_diff < 300:
            print(f"[DUPLICATE] Ignored scan for {name} ({uid})")
            return "ATTENDANCE_DUPLICATE", 200, {'Content-Type': 'text/plain'}

    try:
        wb = openpyxl.load_workbook(ATTENDANCE_FILE)
        ws = wb.active
        ws.append([timestamp_str, uid, name, "PRESENT"])
        wb.save(ATTENDANCE_FILE)

        last_scanned_times[uid] = now
        print(f"[SUCCESS] Attendance Marked: {timestamp_str} | {uid} | {name}")
        return "ATTENDANCE_SUCCESS", 200, {'Content-Type': 'text/plain'}

    except Exception as e:
        print(f"[ERROR] Could not save attendance: {e}")
        return "SERVER_ERROR", 500

@app.route('/add_student', methods=['GET'])
def add_student():
    uid = request.args.get('uid')
    name = request.args.get('name')

    if not uid or not name:
        return "MISSING_PARAMS", 400

    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    try:
        wb = openpyxl.load_workbook(STUDENTS_FILE)
        ws = wb.active

        # Prevent duplicate student UID entries
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] == uid:
                print(f"[EXISTS] Student UID {uid} already exists as {row[1]}")
                return "STUDENT_EXISTS", 200, {'Content-Type': 'text/plain'}

        ws.append([uid, name, now])
        wb.save(STUDENTS_FILE)

        print(f"[REGISTER] New Student Saved: {name} | UID: {uid}")
        return "STUDENT_ADDED", 200, {'Content-Type': 'text/plain'}

    except Exception as e:
        print(f"[ERROR] Could not save student: {e}")
        return "SERVER_ERROR", 500

@app.route('/get_students', methods=['GET'])
def get_students():
    """Sends all registered student records to ESP8266 on boot."""
    try:
        wb = openpyxl.load_workbook(STUDENTS_FILE)
        ws = wb.active

        student_lines = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] and row[1]:
                student_lines.append(f"{row[0]}:{row[1]}")

        response_data = "\n".join(student_lines)
        print(f"[FETCH] Sent {len(student_lines)} students to ESP8266")
        return response_data, 200, {'Content-Type': 'text/plain'}

    except Exception as e:
        print(f"[ERROR] Could not read students file: {e}")
        return "SERVER_ERROR", 500

if __name__ == '__main__':
    initialize_files() # Automatically creates Excel files before starting server
    app.run(host='0.0.0.0', port=5000, threaded=True)